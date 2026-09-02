"""
Shared PowerPoint furniture for everything that prints into the ring binder.

Both the S&F Report deck (sf_report.py) and the Site Feasibility site slides
(deck_export.py) draw their page chrome from here, so a change to the header,
the footer or the table treatment lands on every page of the binder at once.
That is the whole reason this module exists: the two decks used to be separate
codebases in separate languages agreeing by convention, and conventions drift.

Geometry is millimetres throughout. Presentations carry four attributes set by
new_presentation():

    prs._w, prs._h   page size in mm
    prs._m           (left, right, top, bottom) margins in mm
    prs._sz(pt)      point size scaled for the page, so chrome specified at A4
                     comes out the same optical size on an A3 sheet
"""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Mm, Pt
from lxml.etree import fromstring as parse_xml

import brand as BR

HEAD_FONT = BR.HEAD_FONT
BODY_FONT = BR.BODY_FONT

# The en dash is the bullet everywhere. A round bullet reads as a slide deck;
# a dash reads as a document, which is what goes in the binder.
DASH = "–"


# ------------------------------------------------------------------ setup

def new_presentation(size="a4"):
    """A blank deck at the binder page size, with the brand margins applied."""
    w_mm, h_mm = (int(v) for v in BR.PAGE[size])
    prs = Presentation()
    prs.slide_width = Mm(w_mm)
    prs.slide_height = Mm(h_mm)
    prs._w, prs._h = w_mm, h_mm
    prs._m = BR.margins(size)
    # Chrome is specified at A4, the size the deck is actually printed at, then
    # scaled up for an A3 sheet so it reads the same at arm's length.
    chrome = 1.0 if size == "a4" else 1.414
    prs._sz = lambda pt: pt * chrome
    return prs


def blank(prs):
    """A slide with no layout placeholders on it. Everything is positioned."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def content_box(prs):
    """(x, y, width) of the live area, inside the margins."""
    ml, mr, mt, _ = prs._m
    return ml, mt, prs._w - ml - mr


# ------------------------------------------------------------------- text

def _apply(run, size, font, colour, bold, italic, tracking):
    run.font.size = Pt(size)
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = colour
    if tracking:
        # python-pptx has no letter-spacing API. spc is in 100ths of a point.
        run.font._rPr.set("spc", str(int(tracking * 100)))


def _frame(shape, align, valign, spacing):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Mm(0)
    tf.margin_top = tf.margin_bottom = Mm(0)
    if valign is not None:
        tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    return tf


def tb(slide, x, y, w, h, text, size=10, font=BODY_FONT, colour=BR.BODY,
       bold=False, italic=False, align=PP_ALIGN.LEFT, valign=None,
       spacing=1.15, tracking=None):
    """A single-run textbox. The workhorse."""
    box = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    p = _frame(box, align, valign, spacing).paragraphs[0]
    _apply(p.add_run(), size, font, colour, bold, italic, tracking)
    p.runs[0].text = text
    return box


def rich(slide, x, y, w, h, paras, size=10, font=BODY_FONT, colour=BR.BODY,
         bold=False, italic=False, align=PP_ALIGN.LEFT, valign=None,
         spacing=1.35, space_after=0, tracking=None):
    """
    A multi-paragraph textbox.

    `paras` is a list, each entry one of:
        "some text"                         a plain paragraph
        {"text": ..., "bullet": True, ...}  a paragraph with overrides
        {"runs": [(text, {...}), ...]}      mixed formatting on one line

    Per-paragraph and per-run overrides accept size, font, colour, bold,
    italic, tracking, and paragraphs also take bullet, spacing, space_after.
    An empty string gives a blank line, which is how the body copy gets its
    paragraph breaks without fighting paraSpaceAfter.
    """
    box = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = _frame(box, align, valign, spacing)
    base = dict(size=size, font=font, colour=colour, bold=bold, italic=italic,
                tracking=tracking)

    for i, spec in enumerate(paras):
        if isinstance(spec, str):
            spec = {"text": spec}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spec.get("spacing", spacing)
        after = spec.get("space_after", space_after)
        if after:
            p.space_after = Pt(after)
        if spec.get("bullet"):
            _bullet(p)
        pfmt = {k: spec[k] for k in base if k in spec}
        runs = spec.get("runs")
        if runs is None:
            runs = [(spec.get("text", ""), {})]
        for text, over in runs:
            fmt = dict(base, **pfmt)
            fmt.update(over)
            run = p.add_run()
            _apply(run, fmt["size"], fmt["font"], fmt["colour"], fmt["bold"],
                   fmt["italic"], fmt["tracking"])
            run.text = text
    return box


def _bullet(p, char=DASH, hang=4.2):
    """
    Hanging en-dash bullet.

    python-pptx exposes no bullet API. The elements go on pPr by hand; the
    schema orders buFont before buChar, and both after any lnSpc/spcAft that
    the line-spacing and space-after setters put there.
    """
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Mm(hang)))
    pPr.set("indent", str(-Mm(hang)))
    for tag in ("a:buFont", "a:buChar"):
        for existing in pPr.findall(qn(tag)):
            pPr.remove(existing)
    pPr.append(parse_xml('<a:buFont %s typeface="%s"/>'
                         % (nsdecls("a"), BODY_FONT)))
    pPr.append(parse_xml('<a:buChar %s char="%s"/>' % (nsdecls("a"), char)))


# ---------------------------------------------------------------- shapes

def rect(slide, x, y, w, h, fill, line=None, dashed=False, width=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(x), Mm(y), Mm(w), Mm(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(width)
        if dashed:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sh.shadow.inherit = False
    # Autoshapes carry a centred text frame by default, which shows up as a
    # stray empty paragraph if anything later reads the shape.
    sh.text_frame.word_wrap = True
    return sh


def hline(slide, x, y, w, colour=BR.LINE, width=0.75):
    ln = slide.shapes.add_connector(1, Mm(x), Mm(y), Mm(x + w), Mm(y))
    ln.line.color.rgb = colour
    ln.line.width = Pt(width)
    return ln


def full_bleed(slide, colour):
    """Background colour, not a shape. A shape would sit above the footer."""
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = colour


# ------------------------------------------------------------- page chrome

def head(slide, prs, eyebrow, sub, note=None):
    """
    The binder's section head: a small letterspaced green eyebrow over a large
    ink subhead, on white.

    The feasibility slides used to open with a full-width green band. On their
    own that looked fine, but dropped into the binder they read as a different
    document every time the reader turned to one, because no other page has a
    coloured header. The cover is the only page that goes solid green.

    Returns the y in mm where content can start.
    """
    ml, mt, cw = content_box(prs)
    tb(slide, ml, mt, cw, 6, eyebrow.upper(), size=prs._sz(10.5),
       font=HEAD_FONT, colour=BR.GREEN, bold=True, spacing=1.0, tracking=1.5)

    # The subhead sits in a fixed 9mm band with the note pinned below it, so a
    # subhead that wraps prints straight through the note. Nearly every one is
    # two or three words, but the site pages put the full geocoded address here
    # and Auckland's long ones ("12B, Margate Road, Blockhouse Bay, Whau,
    # Auckland, 0600, New Zealand / Aotearoa") run to two lines. Shrink to fit
    # one line, and only if it still will not fit, give it the second line and
    # push everything below it down.
    size = prs._sz(17)
    while size > prs._sz(11) and len(sub) * 0.62 * size * 0.3528 > cw:
        size -= prs._sz(0.5)
    wraps = len(sub) * 0.62 * size * 0.3528 > cw
    drop = size * 1.15 * 0.3528 if wraps else 0
    tb(slide, ml, mt + 5.6, cw, 9 + drop, sub, size=size, font=HEAD_FONT,
       colour=BR.INK, bold=True, spacing=1.15)
    if note:
        tb(slide, ml, mt + 15.5 + drop, cw, 6, note, size=prs._sz(10),
           colour=BR.MUTED)
    return mt + (23.5 if note else 17.5) + drop


def footer(slide, prs, page_number=True):
    """Hairline, brand line left, auto-updating slide number right."""
    ml, mr, _, mb = prs._m
    cw = prs._w - ml - mr
    y = prs._h - mb + 3.5
    hline(slide, ml, y, cw)
    tb(slide, ml, y + 1.6, cw * 0.7, 5, BR.FOOTER_LINE, size=prs._sz(7),
       font=HEAD_FONT, colour=BR.MUTED, tracking=1)
    if page_number:
        slide_number(slide, prs, ml + cw - 20, y + 1.4, 20, 5)


def slide_number(slide, prs, x, y, w, h):
    """
    A real slide-number field, not baked text.

    Optional slides get deleted per job, so a typed number is wrong on every
    page after the first deletion.
    """
    box = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Mm(0)
    tf.margin_top = tf.margin_bottom = Mm(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p._p.append(parse_xml(
        '<a:fld %s id="{1D8BD8AF-4C40-4A3E-9E9A-3E6D4E1B77A1}" '
        'type="slidenum"><a:rPr lang="en-NZ" sz="%d" b="1">'
        '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
        '<a:latin typeface="%s"/></a:rPr><a:t>2</a:t></a:fld>'
        % (nsdecls("a"), int(prs._sz(8) * 100), BR.GREEN_HEX, HEAD_FONT)))
    return box


def banner(slide, prs, text):
    """
    Amber strip across the top of the page.

    Internal only. Amber never survives into a printed client page: it marks
    what Lee fills in or deletes first, and the how-to-use slide says so.
    """
    rect(slide, 0, 0, prs._w, 8.6, BR.NOTE_BG)
    ml, mr, _, _ = prs._m
    tb(slide, ml, 0, prs._w - ml - mr, 8.6, text, size=prs._sz(8.5),
       colour=BR.NOTE_INK, bold=True, italic=True, valign=MSO_ANCHOR.MIDDLE)


def placeholder(slide, prs, x, y, w, h, label, hint=None, label_size=11):
    """
    A dashed box marking where an image goes. Deleted once the image is in.

    label_size drops for a strip of narrow boxes, where the default wraps the
    label and leaves the closing bracket stranded on its own line.
    """
    rect(slide, x, y, w, h, BR.PHBG, BR.PH_EDGE, dashed=True, width=1)
    paras = [{"text": "[ %s ]" % label, "size": prs._sz(label_size),
              "bold": True, "italic": True, "colour": BR.PH_INK}]
    if hint:
        paras.append({"text": hint, "size": prs._sz(8.5), "italic": True,
                      "colour": BR.MUTED})
    rich(slide, x + 3, y, w - 6, h, paras, align=PP_ALIGN.CENTER,
         valign=MSO_ANCHOR.MIDDLE, spacing=1.3)


# ----------------------------------------------------------------- tables

def cell_border(cell, edge, colour, pt):
    """
    python-pptx exposes no border API, so the line elements go in by hand.
    Order matters: lnL, lnR, lnT, lnB must appear in that sequence inside
    tcPr or PowerPoint rejects the table.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tag = "a:ln%s" % edge
    for existing in tcPr.findall(qn(tag)):
        tcPr.remove(existing)
    if pt is None:
        ln = parse_xml('<a:ln%s %s><a:noFill/></a:ln%s>'
                       % (edge, nsdecls("a"), edge))
    else:
        ln = parse_xml(
            '<a:ln%s %s w="%d" cap="flat"><a:solidFill>'
            '<a:srgbClr val="%s"/></a:solidFill></a:ln%s>'
            % (edge, nsdecls("a"), int(pt * 12700), colour, edge))
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    anchor = None
    for later in order[order.index(tag) + 1:]:
        found = tcPr.find(qn(later))
        if found is not None:
            anchor = found
            break
    if anchor is None:
        tcPr.insert(0, ln)
    else:
        anchor.addprevious(ln)


def style_table(table, header=True, pad=1.6):
    """
    The binder's table language: a green rule under the header, a hairline
    between rows, no vertical lines, no banding, white throughout.

    The obvious alternative, a solid green header with zebra body rows, is a
    spreadsheet convention and it fought every other table in the binder.
    Banding also makes a short list of statements look like data to be scanned
    rather than sentences to be read.

    Fills have to be painted explicitly: python-pptx tables inherit Office's
    blue banded default, and setting run colours alone leaves that in place.
    """
    table.first_row = False
    table.horz_banding = False
    n_rows = len(table.rows)
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.margin_left = Mm(1.5)
            cell.margin_right = Mm(3)
            # `pad` is the one lever a long table has that costs nothing to
            # read: on a 15-row findings table the default padding alone is
            # 48mm, so tightening it saves the page before the type has to
            # come down.
            cell.margin_top = Mm(pad)
            cell.margin_bottom = Mm(pad)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BR.WHITE
            cell_border(cell, "L", None, None)
            cell_border(cell, "R", None, None)
            cell_border(cell, "T", None, None)
            if r == 0 and header:
                cell_border(cell, "B", BR.GREEN_HEX, 1.5)
            elif r < n_rows - 1:
                cell_border(cell, "B", BR.LINE_HEX, 0.5)
            else:
                cell_border(cell, "B", None, None)
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.15
                for run in p.runs:
                    if not run.font.name:
                        run.font.name = BODY_FONT


def cell(table, r, c, text, size=10, bold=False, italic=False, colour=BR.INK,
         font=BODY_FONT, tracking=None):
    tc = table.cell(r, c)
    tc.text = ""
    p = tc.text_frame.paragraphs[0]
    run = p.add_run()
    _apply(run, size, font, colour, bold, italic, tracking)
    run.text = text
    return tc


def brand_table(slide, prs, x, y, w, h, col_fracs, headers, rows,
                size=10, head_size=8.5):
    """
    Header row plus body rows in the binder's table treatment.

    A row cell is either a string, or a (text, muted) pair. Muted renders as
    italic grey, which is how every [bracketed] prompt in the template reads:
    visibly not yet written, and visibly Lee's to replace.
    """
    tbl = slide.shapes.add_table(len(rows) + 1, len(headers), Mm(x), Mm(y),
                                 Mm(w), Mm(h)).table
    for i, frac in enumerate(col_fracs):
        tbl.columns[i].width = Mm(w * frac)
    for c, text in enumerate(headers):
        cell(tbl, 0, c, text.upper(), size=prs._sz(head_size), bold=True,
             colour=BR.GREEN, font=HEAD_FONT, tracking=1.5)
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            text, muted = value if isinstance(value, tuple) else (value, False)
            cell(tbl, r, c, text, size=prs._sz(size), italic=muted,
                 colour=BR.MUTED if muted else BR.BODY)
    style_table(tbl)
    return tbl
