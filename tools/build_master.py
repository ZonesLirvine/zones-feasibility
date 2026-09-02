"""
Rebuild the master S&F template from the deck Lee actually sent.

    python tools/build_master.py "<sent-deck.pptx>"

One file, in Drive:

    _Zones/Resources/Templates/SF-Presentation-Template-Landscape-v1.pptx

It is a working document: every page of the current design with {{merge}}
tokens and [bracketed] prompts, on the layouts and theme from the sent deck.
The three YOUR SITE pages are drawn the way a real job draws them, with tokens
where a property's own answers go. The master shows that part of the design
without carrying anybody's data.
Lee opens it to hand-build a report; the generator opens it for its layouts and
drops the pages. One file doing both jobs is the point, because two could
disagree.

Built by the generator rather than maintained by hand, which is the only way
the template and the tool stay the same document.

Edited in place, never duplicated and renamed: a second copy is how the
automation ended up serving an outdated template last time. The mirror under
Operations/Resources/Templates is a separate file rather than a shortcut, so
this writes both.
"""
import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation                                    # noqa: E402

import sf_report                                                 # noqa: E402
import slides as SL                                              # noqa: E402

TEMPLATES = [r"G:\My Drive\_Zones\Resources\Templates",
             r"G:\My Drive\_Zones\Operations\Resources\Templates"]
MASTER = "SF-Presentation-Template-Landscape-v1.pptx"


def main(sent_deck):
    tmp = os.environ.get("TEMP", ".")
    # The sent deck with its pages removed is the layout and theme source. The
    # generator has to build onto that rather than onto the master it is about
    # to replace, or a design change would take two runs to land.
    carrier = os.path.join(tmp, "_carrier.pptx")
    SL.strip_slides(Presentation(sent_deck)).save(carrier)

    built = os.path.join(tmp, "_master.pptx")
    SL.MASTER = carrier
    sf_report.build_template(built)
    SL.MASTER = os.path.join(TEMPLATES[0], MASTER)

    for folder in TEMPLATES:
        if os.path.isdir(folder):
            shutil.copyfile(built, os.path.join(folder, MASTER))
        else:
            print("skipped (not mounted): %s" % folder)
    shutil.copyfile(carrier, SL.TEMPLATE)      # offline cache, already stripped
    os.remove(carrier)
    os.remove(built)

    for folder in TEMPLATES:
        path = os.path.join(folder, MASTER)
        if os.path.exists(path):
            chk = Presentation(path)
            print("%2d pages  %8d bytes  %s"
                  % (len(chk.slides), os.path.getsize(path), path))
    print("cache: %d bytes" % os.path.getsize(SL.TEMPLATE))


if __name__ == "__main__":
    main(sys.argv[1])
